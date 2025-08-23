import unittest
import os
import sys
import tempfile
import shutil
from pathlib import Path
import io
import time

# Add parent directory to path for imports
sys.path.append(os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from codebase_to_text.codebase_to_text import CodebaseToText


class TestCodebaseToText(unittest.TestCase):
    def setUp(self):
        """Set up test environment with temporary folder structure"""
        self.test_folder_path = tempfile.mkdtemp(prefix="test_codebase_")
        
        # Create test folder structure
        self._create_test_structure()
        
        # Output paths for testing
        self.output_txt = os.path.join(self.test_folder_path, "output.txt")
        self.output_docx = os.path.join(self.test_folder_path, "output.docx")

    def _create_test_structure(self):
        """Create a complex test folder structure"""
        base = self.test_folder_path
        
        # Create main files
        with open(os.path.join(base, "main.py"), "w") as f:
            f.write("print('Hello World')")
        
        with open(os.path.join(base, "README.md"), "w") as f:
            f.write("# Test Project\nThis is a test.")
        
        with open(os.path.join(base, "requirements.txt"), "w") as f:
            f.write("requests>=2.25.0\npandas>=1.3.0")
        
        # Create subdirectories
        os.makedirs(os.path.join(base, "src"), exist_ok=True)
        os.makedirs(os.path.join(base, "tests"), exist_ok=True)
        os.makedirs(os.path.join(base, "__pycache__"), exist_ok=True)
        os.makedirs(os.path.join(base, ".git"), exist_ok=True)
        os.makedirs(os.path.join(base, "venv", "lib"), exist_ok=True)
        os.makedirs(os.path.join(base, "logs"), exist_ok=True)
        
        # Create files in subdirectories
        with open(os.path.join(base, "src", "app.py"), "w") as f:
            f.write("def main():\n    pass")
        
        with open(os.path.join(base, "src", "utils.py"), "w") as f:
            f.write("def helper():\n    return True")
        
        with open(os.path.join(base, "tests", "test_app.py"), "w") as f:
            f.write("import unittest\n\nclass TestApp(unittest.TestCase):\n    pass")
        
        # Create files that should be excluded by default
        with open(os.path.join(base, "__pycache__", "app.cpython-39.pyc"), "w") as f:
            f.write("binary content")
        
        with open(os.path.join(base, ".git", "config"), "w") as f:
            f.write("[core]\nrepositoryformatversion = 0")
        
        with open(os.path.join(base, "venv", "lib", "python3.9"), "w") as f:
            f.write("virtual env file")
        
        with open(os.path.join(base, "logs", "app.log"), "w") as f:
            f.write("2023-01-01 10:00:00 INFO Application started")
        
        with open(os.path.join(base, "temp.tmp"), "w") as f:
            f.write("temporary content")
        
        # Create hidden files
        with open(os.path.join(base, ".gitignore"), "w") as f:
            f.write("*.pyc\n__pycache__/\n.env")
        
        with open(os.path.join(base, ".env"), "w") as f:
            f.write("SECRET_KEY=test123")

    def test_basic_functionality(self):
        """Test basic text generation without exclusions"""
        code_to_text = CodebaseToText(
            input_path=self.test_folder_path,
            output_path=self.output_txt,
            output_type="txt",
            verbose=False,
            exclude_hidden=False,
            exclude=[]
        )
        
        text = code_to_text.get_text()
        self.assertIn("Folder Structure", text)
        self.assertIn("File Contents", text)
        self.assertIn("main.py", text)
        self.assertIn("Hello World", text)

    def test_exclude_hidden_files(self):
        """Test exclusion of hidden files"""
        code_to_text = CodebaseToText(
            input_path=self.test_folder_path,
            output_path=self.output_txt,
            output_type="txt",
            verbose=False,
            exclude_hidden=True,
            exclude=[]
        )
        text = code_to_text.get_text()
        self.assertNotIn(".gitignore", text)
        self.assertNotIn(".env", text)
        self.assertIn("main.py", text)  # Regular files should still be included

    def test_exclude_patterns(self):
        """Test pattern-based exclusions"""
        exclude_patterns = ["*.log", "*.tmp", "__pycache__/**", ".git/**"]
        
        code_to_text = CodebaseToText(
            input_path=self.test_folder_path,
            output_path=self.output_txt,
            output_type="txt",
            verbose=False,
            exclude_hidden=False,
            exclude=exclude_patterns
        )
        
        text = code_to_text.get_text()
        
        # Split the text to get only the folder structure section
        folder_structure_section = text.split("File Contents")[0]
        
        # Should exclude log and tmp files from folder structure
        self.assertNotIn("app.log", folder_structure_section)
        self.assertNotIn("temp.tmp", folder_structure_section)
        self.assertNotIn("__pycache__/", folder_structure_section)
        self.assertNotIn(".git/", folder_structure_section)
        
        # Should include normal files in folder structure
        self.assertIn("main.py", folder_structure_section)
        self.assertIn("src/", folder_structure_section)

    def test_exclude_specific_files(self):
        """Test exclusion of specific files"""
        exclude_patterns = ["README.md", "requirements.txt"]
        
        code_to_text = CodebaseToText(
            input_path=self.test_folder_path,
            output_path=self.output_txt,
            output_type="txt",
            verbose=False,
            exclude_hidden=False,
            exclude=exclude_patterns
        )
        
        text = code_to_text.get_text()
          # Should exclude specified files
        self.assertNotIn("README.md", text)
        self.assertNotIn("requirements.txt", text)
        
        # Should include other files
        self.assertIn("main.py", text)

    def test_exclude_directories(self):
        """Test exclusion of entire directories"""
        exclude_patterns = ["venv/", "logs/"]
        
        code_to_text = CodebaseToText(
            input_path=self.test_folder_path,
            output_path=self.output_txt,
            output_type="txt",
            verbose=False,
            exclude_hidden=False,
            exclude=exclude_patterns
        )
        
        text = code_to_text.get_text()
        
        # Split the text to get only the folder structure section
        folder_structure_section = text.split("File Contents")[0]
        
        # Should exclude specified directories from folder structure
        self.assertNotIn("venv/", folder_structure_section)
        self.assertNotIn("logs/", folder_structure_section)
        
        # Should include other directories
        self.assertIn("src/", folder_structure_section)
        self.assertIn("tests/", folder_structure_section)

    def test_exclude_file_creation(self):
        """Test loading exclusion patterns from .exclude file"""
        exclude_file_path = os.path.join(self.test_folder_path, ".exclude")
        
        # Create .exclude file
        with open(exclude_file_path, "w") as f:
            f.write("# This is a comment\n")
            f.write("*.log\n")
            f.write("temp.tmp\n")
            f.write("venv/\n")
            f.write("\n")  # Empty line
        
        code_to_text = CodebaseToText(
            input_path=self.test_folder_path,
            output_path=self.output_txt,
            output_type="txt",
            verbose=False,
            exclude_hidden=False,
            exclude=[]
        )
        
        text = code_to_text.get_text()
        
        # Split the text to get only the folder structure section
        folder_structure_section = text.split("File Contents")[0]
        
        # Should exclude files listed in .exclude file from folder structure
        self.assertNotIn("app.log", folder_structure_section)
        self.assertNotIn("temp.tmp", folder_structure_section)
        self.assertNotIn("venv/", folder_structure_section)

    def test_combined_exclusions(self):
        """Test combination of CLI args and .exclude file"""
        exclude_file_path = os.path.join(self.test_folder_path, ".exclude")
        
        # Create .exclude file
        with open(exclude_file_path, "w") as f:
            f.write("*.log\n")
            f.write("venv/\n")
        
        # Also provide CLI exclusions
        cli_excludes = ["*.tmp", "__pycache__/"]
        
        code_to_text = CodebaseToText(
            input_path=self.test_folder_path,
            output_path=self.output_txt,
            output_type="txt",
            verbose=False,
            exclude_hidden=False,
            exclude=cli_excludes
        )
        
        text = code_to_text.get_text()
        
        # Split the text to get only the folder structure section
        folder_structure_section = text.split("File Contents")[0]
        
        # Should exclude files from both sources from folder structure
        self.assertNotIn("app.log", folder_structure_section)  # From .exclude file
        self.assertNotIn("venv/", folder_structure_section)    # From .exclude file
        self.assertNotIn("temp.tmp", folder_structure_section) # From CLI
        self.assertNotIn("__pycache__/", folder_structure_section) # From CLI

    def test_output_file_generation_txt(self):
        """Test TXT file output generation"""
        code_to_text = CodebaseToText(
            input_path=self.test_folder_path,
            output_path=self.output_txt,
            output_type="txt",
            verbose=False,
            exclude_hidden=False,
            exclude=["*.log", "*.tmp"]
        )
        
        code_to_text.get_file()
        
        # Check if output file was created
        self.assertTrue(os.path.exists(self.output_txt))
        
        # Check content
        with open(self.output_txt, "r", encoding="utf-8") as f:
            content = f.read()
            self.assertIn("Folder Structure", content)
            self.assertIn("main.py", content)

    def test_output_file_generation_docx(self):
        """Test DOCX file output generation"""
        code_to_text = CodebaseToText(
            input_path=self.test_folder_path,
            output_path=self.output_docx,
            output_type="docx",
            verbose=False,
            exclude_hidden=False,
            exclude=["*.log", "*.tmp"]
        )
        
        code_to_text.get_file()
        
        # Check if output file was created
        self.assertTrue(os.path.exists(self.output_docx))

    def test_verbose_mode(self):
        """Test verbose output mode"""
        import io
        import sys
        
        # Capture stdout
        captured_output = io.StringIO()
        sys.stdout = captured_output
        
        try:
            code_to_text = CodebaseToText(
                input_path=self.test_folder_path,
                output_path=self.output_txt,
                output_type="txt",
                verbose=True,
                exclude_hidden=False,
                exclude=["*.log"]
            )
            
            code_to_text.get_file()
            
            # Get the output
            output = captured_output.getvalue()
            
            # Should contain verbose messages
            self.assertIn("Active exclusion patterns", output)
            self.assertIn("Processing:", output)
        finally:
            # Restore stdout
            sys.stdout = sys.__stdout__

    def test_invalid_output_type(self):
        """Test error handling for invalid output type"""
        with self.assertRaises(ValueError):
            code_to_text = CodebaseToText(
                input_path=self.test_folder_path,
                output_path="output.xyz",
                output_type="xyz",
                verbose=False,
                exclude_hidden=False,
                exclude=[]
            )
            code_to_text.get_file()    
            
    def test_exclusion_count_tracking(self):
        """Test that exclusion counting works correctly"""
        code_to_text = CodebaseToText(
            input_path=self.test_folder_path,
            output_path=self.output_txt,
            output_type="txt",
            verbose=True,  # Need verbose mode for this test to work properly
            exclude_hidden=False,
            exclude=["*.log", "*.tmp", "__pycache__/**"]
        )
        
        # Generate text to trigger exclusion counting
        code_to_text.get_text()
        
        # Should have excluded some files
        self.assertGreater(code_to_text.excluded_files_count, 0)

    def tearDown(self):
        """Clean up test environment"""
        if os.path.exists(self.test_folder_path):
            shutil.rmtree(self.test_folder_path)


class TestPatternMatching(unittest.TestCase):
    """Test exclusion pattern matching specifically"""
    
    def setUp(self):
        self.test_folder_path = tempfile.mkdtemp(prefix="test_patterns_")
        self.code_to_text = CodebaseToText(
            input_path=self.test_folder_path,
            output_path="dummy.txt",
            output_type="txt",
            exclude=[]
        )

    def test_wildcard_patterns(self):
        """Test wildcard pattern matching"""
        self.code_to_text.exclude_patterns = {"*.py", "*.log"}
        
        # Should match
        self.assertTrue(self.code_to_text._should_exclude("test.py", self.test_folder_path))
        self.assertTrue(self.code_to_text._should_exclude("app.log", self.test_folder_path))
        
        # Should not match
        self.assertFalse(self.code_to_text._should_exclude("test.txt", self.test_folder_path))
        self.assertFalse(self.code_to_text._should_exclude("README.md", self.test_folder_path))

    def test_directory_patterns(self):
        """Test directory pattern matching"""
        self.code_to_text.exclude_patterns = {"__pycache__/", "node_modules/"}
        
        # Create test directories
        pycache_dir = os.path.join(self.test_folder_path, "__pycache__")
        os.makedirs(pycache_dir, exist_ok=True)
        
        # Should match directories
        self.assertTrue(self.code_to_text._should_exclude(pycache_dir, self.test_folder_path))

    def test_recursive_patterns(self):
        """Test recursive wildcard patterns"""
        self.code_to_text.exclude_patterns = {"**/__pycache__/**", "**/node_modules/**"}
        
        # Create nested test structure
        nested_pycache = os.path.join(self.test_folder_path, "src", "utils", "__pycache__", "file.pyc")
        os.makedirs(os.path.dirname(nested_pycache), exist_ok=True)
        
        # Should match nested paths
        self.assertTrue(self.code_to_text._should_exclude(nested_pycache, self.test_folder_path))

    def tearDown(self):
        """Clean up test environment"""
        if os.path.exists(self.test_folder_path):
            shutil.rmtree(self.test_folder_path)


class TestDocxImage(unittest.TestCase):
    def test_docx_with_image(self):
        import tempfile
        import os
        import base64
        from docx import Document
        # Create a temporary directory
        with tempfile.TemporaryDirectory() as temp_dir:
            # Write a dummy 1x1 PNG image
            img_path = os.path.join(temp_dir, "dummy.png")
            png_data = "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABAQAAAAA3bvkkAAAAC0lEQVQIW2NgYAAAAAMAAWgmWQ0AAAAASUVORK5CYII="
            with open(img_path, "wb") as f:
                f.write(base64.b64decode(png_data))
            
            # Create a new DOCX document and add a picture
            doc = Document()
            doc.add_paragraph("Testing DOCX image inclusion.")
            doc.add_picture(img_path)
            
            # Save the document
            doc_path = os.path.join(temp_dir, "test.docx")
            doc.save(doc_path)
            
            # Reload the document and assert that it contains an inline image
            new_doc = Document(doc_path)
            self.assertGreater(len(new_doc.inline_shapes), 0, "Document should contain at least one inline image.")


class TestSecurityPathValidation(unittest.TestCase):
    """Test path validation security features"""
    
    def setUp(self):
        """Set up test environment with temporary folder structure"""
        self.test_folder_path = tempfile.mkdtemp(prefix="test_security_")
        self.external_folder_path = tempfile.mkdtemp(prefix="test_external_")
        
        # Create test files
        with open(os.path.join(self.test_folder_path, "safe_file.txt"), "w") as f:
            f.write("This is a safe file")
        
        # Create external file that should not be accessible
        with open(os.path.join(self.external_folder_path, "secret.txt"), "w") as f:
            f.write("This should not be accessible")
        
        self.output_path = os.path.join(self.test_folder_path, "output.txt")

    def test_validate_file_path_safe_files(self):
        """Test that legitimate files pass validation"""
        converter = CodebaseToText(
            input_path=self.test_folder_path,
            output_path=self.output_path,
            output_type="txt"
        )
        
        safe_file = os.path.join(self.test_folder_path, "safe_file.txt")
        self.assertTrue(converter._validate_file_path(safe_file, self.test_folder_path))

    def test_validate_file_path_traversal_attack(self):
        """Test that directory traversal attempts are blocked"""
        converter = CodebaseToText(
            input_path=self.test_folder_path,
            output_path=self.output_path,
            output_type="txt"
        )
        
        # Test various traversal patterns
        traversal_paths = [
            "../secret.txt",
            "../../secret.txt",
            "../../../etc/passwd",
            "..\\..\\windows\\system32\\config\\sam",  # Windows style
            os.path.join(self.test_folder_path, "../secret.txt"),
        ]
        
        for malicious_path in traversal_paths:
            with self.subTest(path=malicious_path):
                result = converter._validate_file_path(malicious_path, self.test_folder_path)
                self.assertFalse(result, f"Should reject traversal path: {malicious_path}")

    def test_validate_file_path_absolute_paths(self):
        """Test that absolute paths outside base directory are blocked"""
        converter = CodebaseToText(
            input_path=self.test_folder_path,
            output_path=self.output_path,
            output_type="txt"
        )
        
        # Test absolute paths outside the base directory
        external_file = os.path.join(self.external_folder_path, "secret.txt")
        self.assertFalse(converter._validate_file_path(external_file, self.test_folder_path))
        
        # Test system paths
        system_paths = ["/etc/passwd", "/windows/system32/config/sam", "C:\\Windows\\System32"]
        for sys_path in system_paths:
            with self.subTest(path=sys_path):
                result = converter._validate_file_path(sys_path, self.test_folder_path)
                self.assertFalse(result, f"Should reject system path: {sys_path}")

    def test_validate_file_path_symlinks(self):
        """Test that symlinks pointing outside base directory are blocked"""
        converter = CodebaseToText(
            input_path=self.test_folder_path,
            output_path=self.output_path,
            output_type="txt"
        )
        
        # Create a symlink pointing to external file
        symlink_path = os.path.join(self.test_folder_path, "malicious_symlink.txt")
        external_target = os.path.join(self.external_folder_path, "secret.txt")
        
        try:
            os.symlink(external_target, symlink_path)
            # Symlink should be rejected because it points outside base directory
            result = converter._validate_file_path(symlink_path, self.test_folder_path)
            self.assertFalse(result, "Should reject symlink pointing outside base directory")
        except OSError:
            # Skip test if symlinks not supported on this system
            self.skipTest("Symlinks not supported on this system")

    def test_validate_file_path_symlinks_internal(self):
        """Test that symlinks pointing within base directory are allowed"""
        converter = CodebaseToText(
            input_path=self.test_folder_path,
            output_path=self.output_path,
            output_type="txt"
        )
        
        # Create a symlink pointing to internal file
        symlink_path = os.path.join(self.test_folder_path, "internal_symlink.txt")
        internal_target = os.path.join(self.test_folder_path, "safe_file.txt")
        
        try:
            os.symlink(internal_target, symlink_path)
            # Internal symlink should be allowed
            result = converter._validate_file_path(symlink_path, self.test_folder_path)
            self.assertTrue(result, "Should allow symlink pointing within base directory")
        except OSError:
            # Skip test if symlinks not supported on this system
            self.skipTest("Symlinks not supported on this system")

    def test_security_verbose_logging(self):
        """Test that security violations are logged in verbose mode"""
        # Capture stdout
        captured_output = io.StringIO()
        original_stdout = sys.stdout
        sys.stdout = captured_output
        
        try:
            converter = CodebaseToText(
                input_path=self.test_folder_path,
                output_path=self.output_path,
                output_type="txt",
                verbose=True
            )
            
            # Test validation with malicious path
            malicious_path = "../../../etc/passwd"
            converter._validate_file_path(malicious_path, self.test_folder_path)
            
            # Get the output
            output = captured_output.getvalue()
            
            # Should contain security warning
            self.assertIn("SECURITY:", output)
            self.assertIn("Rejected potentially unsafe path", output)
            
        finally:
            # Restore stdout
            sys.stdout = original_stdout

    def test_process_single_file_security_integration(self):
        """Test that _process_single_file properly integrates security validation"""
        # Create a mock file structure that simulates a traversal attack
        test_dir = tempfile.mkdtemp(prefix="test_integration_")
        
        try:
            # Create legitimate file
            safe_file = os.path.join(test_dir, "safe.txt")
            with open(safe_file, "w") as f:
                f.write("Safe content")
            
            converter = CodebaseToText(
                input_path=test_dir,
                output_path="dummy.txt",
                output_type="txt",
                verbose=True
            )
            
            # Test processing safe file (should work)
            result = converter._process_single_file("safe.txt", test_dir, test_dir)
            self.assertIsNotNone(result)
            self.assertIn("Safe content", result)
            
            # Test processing file with traversal path (should be blocked)
            # Note: We can't actually create a file with ".." in the name on most systems,
            # but we can test the validation logic directly
            self.assertFalse(converter._validate_file_path("../../../etc/passwd", test_dir))
            
        finally:
            shutil.rmtree(test_dir)

    def test_full_conversion_with_security(self):
        """Test that full conversion process respects security measures"""
        converter = CodebaseToText(
            input_path=self.test_folder_path,
            output_path=self.output_path,
            output_type="txt"
        )
        
        # This should complete successfully with only safe files
        text = converter.get_text()
        
        # Should contain safe content
        self.assertIn("safe_file.txt", text)
        self.assertIn("This is a safe file", text)
        
        # Generate the actual file
        converter.get_file()
        self.assertTrue(os.path.exists(self.output_path))

    def test_performance_impact(self):
        """Test that security validation has minimal performance impact"""
        # Create a larger test structure
        large_test_dir = tempfile.mkdtemp(prefix="test_performance_")
        
        try:
            # Create multiple files
            for i in range(100):
                with open(os.path.join(large_test_dir, f"file_{i}.txt"), "w") as f:
                    f.write(f"Content of file {i}")
            
            # Test with security (current implementation)
            start_time = time.time()
            converter_secure = CodebaseToText(
                input_path=large_test_dir,
                output_path="dummy1.txt",
                output_type="txt"
            )
            text1 = converter_secure.get_text()
            secure_time = time.time() - start_time
            
            # For performance comparison, we can't easily disable security,
            # but we can verify the time is reasonable (< 5 seconds for 100 files)
            self.assertLess(secure_time, 5.0, "Security validation should not significantly impact performance")
            
            # Verify content is still correct
            self.assertIn("file_0.txt", text1)
            self.assertIn("Content of file 0", text1)
            
        finally:
            shutil.rmtree(large_test_dir)

    def test_error_handling_in_validation(self):
        """Test error handling in path validation"""
        converter = CodebaseToText(
            input_path=self.test_folder_path,
            output_path=self.output_path,
            output_type="txt",
            verbose=True
        )
        
        # Test with invalid/problematic paths
        invalid_paths = [
            "",  # Empty string
            None,  # This might cause TypeError, but should be handled
            "\x00invalid",  # Null byte
        ]
        
        for invalid_path in invalid_paths:
            with self.subTest(path=invalid_path):
                try:
                    result = converter._validate_file_path(invalid_path, self.test_folder_path)
                    # Should either return False or handle the error gracefully
                    self.assertFalse(result, f"Should reject invalid path: {invalid_path}")
                except (TypeError, ValueError):
                    # These exceptions are acceptable for truly invalid inputs
                    pass

    def tearDown(self):
        """Clean up test environment"""
        if os.path.exists(self.test_folder_path):
            shutil.rmtree(self.test_folder_path)
        if os.path.exists(self.external_folder_path):
            shutil.rmtree(self.external_folder_path)


class TestSecurityIntegration(unittest.TestCase):
    """Integration tests for security features"""
    
    def test_malicious_repo_simulation(self):
        """Simulate a malicious repository with traversal attempts"""
        # This test simulates what could happen with a malicious Git repository
        test_dir = tempfile.mkdtemp(prefix="test_malicious_")
        external_dir = tempfile.mkdtemp(prefix="test_external_")
        
        try:
            # Create legitimate files
            with open(os.path.join(test_dir, "README.md"), "w") as f:
                f.write("# Fake innocent project")
            
            with open(os.path.join(test_dir, "main.py"), "w") as f:
                f.write("print('Hello World')")
            
            # Create external sensitive file
            with open(os.path.join(external_dir, "secrets.txt"), "w") as f:
                f.write("SECRET_API_KEY=abc123")
            
            # Test conversion with verbose mode to catch security messages
            captured_output = io.StringIO()
            original_stdout = sys.stdout
            sys.stdout = captured_output
            
            try:
                converter = CodebaseToText(
                    input_path=test_dir,
                    output_path=os.path.join(test_dir, "output.txt"),
                    output_type="txt",
                    verbose=True
                )
                
                # This should work without accessing external files
                text = converter.get_text()
                
                # Should contain legitimate files
                self.assertIn("README.md", text)
                self.assertIn("main.py", text)
                self.assertIn("Hello World", text)
                
                # Should NOT contain external secrets
                self.assertNotIn("SECRET_API_KEY", text)
                self.assertNotIn("abc123", text)
                
                output = captured_output.getvalue()
                
            finally:
                sys.stdout = original_stdout
                
        finally:
            shutil.rmtree(test_dir)
            shutil.rmtree(external_dir)

if __name__ == "__main__":
    # Run specific test class or all tests
    unittest.main(verbosity=2)